#!/usr/bin/env python3
"""Build Kindle DOCX + paperback interior PDF from one Markdown manuscript.

Usage:
  python3 tools/book_build/build.py books/<slug> [--out books/<slug>/release/v<version>]

Single source: books/<slug>/manuscript/*.md (read in filename order) + books/<slug>/book.json.
The manuscript is never modified. Target differences live here:

  common   footnote labels made unique per file; stable ids on parts/chapters/moments;
           placeholder lines (REVIEW-LINK, @HANDLE) filled from book.json or removed.
  kindle   pandoc -> DOCX with a generated reference.docx (Heading 1 starts a new page,
           no headers/footers/page numbers); Contents + Find Your Moment become
           hyperlinks to heading bookmarks; front-matter rules become page breaks.
  print    pandoc -> HTML -> WeasyPrint PDF at the trim size in book.json: mirrored
           margins, running heads, page numbers, contents/index page numbers computed
           from the final layout (target-counter), per-page footnotes, printable cards
           kept whole, headings kept with following text, embedded OFL fonts.
"""
import argparse, glob, html, json, os, re, shutil, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))


def pandoc_bin():
    try:
        import pypandoc
        return pypandoc.get_pandoc_path()
    except Exception:
        return shutil.which("pandoc") or sys.exit("pandoc not found: pip install pypandoc_binary")


def slugify(text):
    text = re.sub(r"[*_`\"'“”‘’!?.,:;()]", "", text.lower())
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


PAGEBREAK = ('\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'
             '```{=html}\n<div class="pagebreak"></div>\n```\n')
CARD_TITLE = re.compile(r"^\*\*[A-Z0-9][A-Z0-9 :'’&\-]+\*\*$")


# ---------------------------------------------------------------- common source

def load_sources(book_dir, cfg):
    files = sorted(glob.glob(os.path.join(book_dir, "manuscript", "*.md")))
    parts = []
    for f in files:
        stem = os.path.basename(f).split("-")[0]
        s = open(f, encoding="utf-8").read()
        s = re.sub(r"\[\^([A-Za-z0-9]+)\]", lambda m: f"[^f{stem}{m.group(1)}]", s)
        parts.append((os.path.basename(f), s))
    return parts


def list_blank_lines(s):
    """A list directly under a text line needs a blank line to parse as a list (cards/forms)."""
    L = s.split("\n"); out = []
    for i, line in enumerate(L):
        if i and re.match(r"(- |\d+\. )", line) and L[i - 1].strip() and not re.match(r"(\s|- |\d+\. |#)", L[i - 1]):
            out.append("")
        out.append(line)
    return "\n".join(out)


def escape_blanks(s):
    """Write-in blanks (___) are literal underscores, not Markdown emphasis."""
    return re.sub(r"_{3,}", lambda m: "\\_" * len(m.group(0)), s)


def fill_placeholders(s, cfg):
    review, handle = cfg["links"].get("review_url"), cfg["links"].get("instagram_handle")
    if review:
        s = s.replace("(REVIEW-LINK)", f"({review})")
    else:
        s = re.sub(r"^\*\*\[Leave a review on Amazon\]\(REVIEW-LINK\)\*\*\n\n?", "", s, flags=re.M)
    if handle:
        s = s.replace("@HANDLE", "@" + handle.lstrip("@"))
    else:
        s = re.sub(r"^- Instagram: \*\*@HANDLE\*\*\n", "", s, flags=re.M)
    return s


def add_ids(s):
    s = re.sub(r"^# (Part (\d+): .+)$", lambda m: f"# {m.group(1)} {{#part-{m.group(2)} .part}}", s, flags=re.M)
    s = re.sub(r"^# (Chapter (\d+): .+)$", lambda m: f"# {m.group(1)} {{#chapter-{m.group(2)} .chapter}}", s, flags=re.M)
    s = re.sub(r"^## (Moment (\d+): .+)$", lambda m: f"## {m.group(1)} {{#moment-{m.group(2)} .moment}}", s, flags=re.M)

    def other_h1(m):
        t = m.group(1)
        if "{#" in t:
            return m.group(0)
        return f"# {t} {{#{slugify(t)}}}"
    s = re.sub(r"^# (?!Part \d|Chapter \d)(.+)$", other_h1, s, flags=re.M)
    return s


def heading_ids(src):
    ids = {}
    for m in re.finditer(r"^# (.+?) \{#([a-z0-9-]+)", src, re.M):
        ids[m.group(1).strip()] = m.group(2)
    return ids


def split_front(front):
    """front matter = title block | copyright | Find Your Moment | Contents (separated by ---)."""
    blocks = re.split(r"^---\s*$", front, flags=re.M)
    if len(blocks) != 4:
        sys.exit(f"front matter: expected 4 blocks separated by ---, got {len(blocks)}")
    return [b.strip("\n") for b in blocks]


def title_block(cfg):
    return (f'::: {{.titlepage custom-style="Title"}}\n{cfg["title"]}\n:::\n\n'
            f'::: {{.subtitle custom-style="Subtitle"}}\n{cfg["subtitle"]}\n:::\n\n'
            f'::: {{.subtitle2 custom-style="Subtitle 2"}}\n{cfg["subtitle_line2"]}\n:::\n\n'
            f'::: {{.series custom-style="Series"}}\n{cfg["series_line"]}\n:::\n\n'
            f'::: {{.author custom-style="Author"}}\n{cfg["author"]}\n:::\n')


def check_title_block(block, cfg):
    need = [cfg["title"], cfg["subtitle"], cfg["subtitle_line2"], cfg["series_line"], cfg["author"]]
    missing = [n for n in need if n not in block]
    if missing:
        sys.exit(f"book.json disagrees with the manuscript title block: {missing}")


def linkify_lists(fym, contents, ids, target):
    # Find Your Moment: list items in order -> moment-1..30
    n = 0
    def fym_item(m):
        nonlocal n
        n += 1
        return f"- [{m.group(1)}](#moment-{n})"
    fym = re.sub(r"^- (.+)$", fym_item, fym, flags=re.M)
    if n != 30:
        sys.exit(f"Find Your Moment has {n} entries, expected 30")

    def toc_item(m):
        indent, bold, text = m.group(1), m.group(2), m.group(3)
        key = text.strip("*")
        if key not in ids:
            sys.exit(f"Contents entry has no matching heading: {key!r}")
        link = f"[{key}](#{ids[key]})"
        return f"{indent}- {'**' + link + '**' if bold else link}"
    contents = re.sub(r"^(\s*)- (\*\*)?(.+?)(?:\*\*)?$", toc_item, contents, flags=re.M)
    cls = " .print-refs" if target == "print" else ""
    fym = fym.replace("# Find Your Moment {#find-your-moment}", "# Find Your Moment {#find-your-moment .fym" + cls + "}")
    contents = contents.replace("# Contents {#contents}", "# Contents {#contents .toc" + cls + "}")
    return fym, contents


def card_linebreaks(s):
    """Printable cards are forms: each source line is its own line (hard break), both targets."""
    out, lines, i = [], s.split("\n"), 0
    while i < len(lines):
        if lines[i].strip() == "---" and i + 2 < len(lines) and CARD_TITLE.match(lines[i + 2].strip()):
            j = i + 1
            while j < len(lines) and lines[j].strip() != "---":
                j += 1
            block = lines[i + 1:j]
            for k in range(len(block) - 1):
                cur, nxt = block[k], block[k + 1]
                if cur.strip() and nxt.strip() and not nxt.lstrip()[:1].islower() \
                        and not nxt.lstrip().startswith(("- ", "1.", "2.", "3.", "4.", "5.", "6.")) \
                        and not cur.lstrip().startswith(("- ",)) and not re.match(r"\s*\d+\.", cur):
                    block[k] = cur.rstrip() + "\\"
            out += [lines[i]] + block + [lines[j] if j < len(lines) else ""]
            i = j + 1
            continue
        out.append(lines[i]); i += 1
    return "\n".join(out)


def wrap_cards(s):
    """print only: --- / **CARD TITLE** ... / --- becomes a boxed ::: card div."""
    out, lines, i = [], s.split("\n"), 0
    while i < len(lines):
        if lines[i].strip() == "---" and i + 2 < len(lines) and CARD_TITLE.match(lines[i + 2].strip()):
            j = i + 1
            while j < len(lines) and lines[j].strip() != "---":
                j += 1
            out += ["::: card"] + lines[i + 1:j] + [":::"]
            i = j + 1
            continue
        out.append(lines[i]); i += 1
    return "\n".join(out)


def assemble(book_dir, cfg, target):
    parts = load_sources(book_dir, cfg)
    body = []
    for name, s in parts:
        s = list_blank_lines(escape_blanks(fill_placeholders(s, cfg)))
        body.append((name, s))
    joined = "\n\n".join(s for n, s in body[1:])
    joined = card_linebreaks(add_ids(joined))
    ids = heading_ids(joined)
    title, copyright_, fym, contents = split_front(body[0][1])
    check_title_block(title, cfg)
    fym, contents = linkify_lists(add_ids(fym), add_ids(contents), ids, target)
    if target == "print":
        joined = wrap_cards(joined)
        front = (title_block(cfg) + PAGEBREAK + '::: copyright\n' + copyright_ + '\n:::\n'
                 + PAGEBREAK + fym + "\n\n" + contents + "\n")
    else:
        # Kindle: VML horizontal rules convert unreliably -> centred text separator
        joined = re.sub(r"^---\s*$", '::: {custom-style="Separator"}\n· · ·\n:::', joined, flags=re.M)
        front = (title_block(cfg) + PAGEBREAK + copyright_ + "\n" + PAGEBREAK
                 + fym + "\n\n" + contents + "\n")
    leftovers = re.findall(r"REVIEW-LINK|@HANDLE|\bTODO\b|\bTK\b|<!--", front + joined)
    if leftovers:
        sys.exit(f"unresolved placeholders/comments: {leftovers}")
    return front + "\n\n" + joined + "\n"


# ---------------------------------------------------------------- kindle docx

def make_reference_docx(path):
    import docx
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    subprocess.run([pandoc_bin(), "-o", path, "--print-default-data-file", "reference.docx"], check=True)
    d = docx.Document(path)
    st = d.styles

    def para(name, size=None, bold=None, italic=None, center=False, before=None, after=None,
             page_break=False, keep_next=False, font=None):
        try:
            s = st[name]
        except KeyError:
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH); s.base_style = st["Normal"]
        f = s.font
        if size: f.size = Pt(size)
        if bold is not None: f.bold = bold
        if italic is not None: f.italic = italic
        if font: f.name = font
        pf = s.paragraph_format
        if center: pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if before is not None: pf.space_before = Pt(before)
        if after is not None: pf.space_after = Pt(after)
        pf.page_break_before = page_break
        if keep_next: pf.keep_with_next = True
        return s

    para("Normal", font="Georgia", size=11)
    for n in ("Body Text", "First Paragraph", "Compact"):
        s = para(n, after=8)
        s.paragraph_format.first_line_indent = Inches(0)
        s.paragraph_format.line_spacing = 1.15
    para("Heading 1", size=20, bold=True, center=True, before=36, after=18, page_break=True, keep_next=True)
    para("Heading 2", size=15, bold=True, before=18, after=8, keep_next=True)
    para("Heading 3", size=12.5, bold=True, before=12, after=6, keep_next=True)
    for h in ("Heading 1", "Heading 2", "Heading 3"):
        st[h].font.italic = False
    para("Title", size=28, bold=True, center=True, before=120, after=12)
    para("Subtitle", size=16, italic=False, center=True, after=6)
    para("Subtitle 2", size=13, italic=True, center=True, after=36)
    para("Series", size=11, italic=True, center=True, after=48)
    para("Author", size=16, bold=True, center=True)
    para("Block Text", italic=False, before=6, after=10)
    para("Separator", center=True, before=10, after=10)
    # KDP reflowable: pandoc's default reference has no header/footer parts; never create any.
    d.save(path)


def build_docx(md, out_path, cfg, work):
    ref = os.path.join(work, "reference.docx")
    make_reference_docx(ref)
    src = os.path.join(work, "kindle.md")
    open(src, "w", encoding="utf-8").write(md)
    subprocess.run([pandoc_bin(), src, "-f", "markdown+smart", "-t", "docx", "--reference-doc", ref,
                    "-o", out_path], check=True)
    import docx
    d = docx.Document(out_path)
    d.core_properties.title = cfg["title"]
    d.core_properties.author = cfg["author"]
    d.core_properties.subject = cfg["subtitle"]
    d.core_properties.language = cfg["language"]
    d.save(out_path)


# ---------------------------------------------------------------- print pdf

PRINT_CSS = """
@font-face { font-family: "Book Serif"; src: url("{F}/SourceSerif4-Regular.ttf"); font-weight: 400; font-style: normal; }
@font-face { font-family: "Book Serif"; src: url("{F}/SourceSerif4-Italic.ttf"); font-weight: 400; font-style: italic; }
@font-face { font-family: "Book Serif"; src: url("{F}/SourceSerif4-SemiBold.ttf"); font-weight: 700; font-style: normal; }
@font-face { font-family: "Book Serif"; src: url("{F}/SourceSerif4-SemiBoldItalic.ttf"); font-weight: 700; font-style: italic; }
@font-face { font-family: "Book Sans"; src: url("{F}/Nunito-SemiBold.ttf"); font-weight: 600; }
@font-face { font-family: "Book Sans"; src: url("{F}/Nunito-Bold.ttf"); font-weight: 700; }
@font-face { font-family: "Book Sans"; src: url("{F}/Nunito-ExtraBold.ttf"); font-weight: 800; }

@page {
  size: {W}in {H}in;
  margin: {MT}in {MO}in {MB}in {MI}in;
  @bottom-center { content: none; }
}
@page :left {
  margin-left: {MO}in; margin-right: {MI}in;
  @top-left { content: string(vhead, first-except); font: 600 7.5pt "Book Sans"; letter-spacing: 0.08em; text-transform: uppercase; color: #333; vertical-align: bottom; padding-bottom: 6pt; }
  @bottom-left { content: counter(page); font: 400 9pt "Book Serif"; vertical-align: top; padding-top: 10pt; }
}
@page :right {
  margin-left: {MI}in; margin-right: {MO}in;
  @top-right { content: string(chaptitle, first-except); font: 600 7.5pt "Book Sans"; letter-spacing: 0.08em; text-transform: uppercase; color: #333; vertical-align: bottom; padding-bottom: 6pt; }
  @bottom-right { content: counter(page); font: 400 9pt "Book Serif"; vertical-align: top; padding-top: 10pt; }
}
@page plain { @top-left { content: none; } @top-right { content: none; } @bottom-left { content: none; } @bottom-right { content: none; } }

html { font-family: "Book Serif"; font-size: 10.5pt; line-height: 1.45; color: #000; }
body { margin: 0; }
p { margin: 0 0 0.55em 0; text-align: justify; hyphens: auto; hyphenate-character: "-"; hyphenate-limit-chars: 6 3 3; orphans: 2; widows: 2; }
li p { margin: 0 0 0.2em 0; }
ul, ol { margin: 0.2em 0 0.7em 0; padding-left: 1.3em; }
ol { padding-left: 1.75em; }
li { margin-bottom: 0.25em; text-align: left; hyphens: manual; orphans: 2; widows: 2; }
strong { font-weight: 700; }
a { color: inherit; text-decoration: none; }

h1, h2, h3 { font-family: "Book Sans"; hyphens: manual; text-align: left; break-after: avoid; page-break-after: avoid; break-inside: avoid; }
h1 { font-weight: 800; font-size: 21pt; line-height: 1.15; margin: 0.9in 0 0.35in 0; break-before: page; string-set: vhead "{TITLE}", chaptitle attr(data-rh); }
h2 { font-weight: 700; font-size: 13.5pt; line-height: 1.2; margin: 1.3em 0 0.45em 0; }
h3 { font-weight: 700; font-size: 11pt; margin: 1em 0 0.35em 0; }
h1.chapter .ch-num { display: block; font-weight: 700; font-size: 10pt; letter-spacing: 0.18em; text-transform: uppercase; margin-bottom: 0.5em; }
h1.chapter .ch-title { display: block; }
div.partpage { page: plain; }
h1.part { margin-top: 2.4in; text-align: center; font-size: 22pt; string-set: none; }
h1.part .ch-num { display: block; font-size: 11pt; letter-spacing: 0.2em; text-transform: uppercase; margin-bottom: 0.6em; font-weight: 700; }
h1.part + p { text-align: center; font-style: italic; }
h2.moment { font-size: 12.5pt; margin-top: 1.4em; }
h2.moment + p { font-style: normal; }

p.subhead { break-after: avoid; page-break-after: avoid; margin-bottom: 0.4em; }
p.label { margin-bottom: 0.2em; break-after: avoid; page-break-after: avoid; }
p.script { font-weight: 700; margin-left: 0.18in; text-align: left; hyphens: manual; break-inside: avoid; }
p.stage { margin-left: 0.18in; margin-bottom: 0.2em; break-after: avoid; }
blockquote { margin: 0.6em 0 0.9em 0.18in; padding: 0; break-inside: avoid; }
blockquote p { text-align: left; hyphens: manual; }
hr { border: none; margin: 0.9em 0; text-align: center; height: 1em; }
hr::after { content: "\\00B7\\2003\\00B7\\2003\\00B7"; font-size: 10pt; }

div.card { border: 1.2pt solid #000; border-radius: 6pt; padding: 0.16in 0.2in 0.08in; margin: 1em 0; break-inside: avoid; page-break-inside: avoid; }
div.card p, div.card li { text-align: left; hyphens: manual; }
div.card > p:first-child { font-family: "Book Sans"; font-weight: 800; letter-spacing: 0.04em; margin-bottom: 0.6em; }

h1.runon { break-before: auto; page-break-before: auto; margin-top: 0.45in; font-size: 16pt; }
div.pagebreak { break-after: page; page-break-after: always; height: 0; }
div[data-custom-style] { text-align: center; hyphens: manual; page: plain; }
div[data-custom-style] p { text-align: center; hyphens: manual; }
div[data-custom-style="Title"] p { font-family: "Book Sans"; font-weight: 800; font-size: 30pt; line-height: 1.1; margin-top: 1.9in; margin-bottom: 0.25in; }
div[data-custom-style="Subtitle"] p { font-family: "Book Sans"; font-weight: 700; font-size: 14pt; line-height: 1.3; margin-bottom: 0.12in; }
div[data-custom-style="Subtitle 2"] p { font-style: italic; font-size: 11pt; margin-bottom: 0.6in; }
div[data-custom-style="Series"] p { font-style: italic; font-size: 10pt; margin-bottom: 1.4in; }
div[data-custom-style="Author"] p { font-family: "Book Sans"; font-weight: 700; font-size: 14pt; letter-spacing: 0.12em; text-transform: uppercase; }
div.copyright { page: plain; font-size: 8.5pt; line-height: 1.4; padding-top: 3.2in; }
div.copyright p { text-align: left; hyphens: manual; }

div.frontlist { page: plain; }
h1.fym, h1.toc { margin-top: 0.4in; string-set: none; }
h1.fym ~ p { text-align: left; }
.print-refs + p { font-style: italic; }
ul.refs { list-style: none; padding-left: 0; margin-top: 0.1em; }
ul.refs li { margin-bottom: 0.18em; }
ul.refs li a::after { content: leader(". ") target-counter(attr(href), page); }
ul.refs ul { list-style: none; padding-left: 1.1em; margin: 0.15em 0 0.3em 0; }
.toc-block ul.refs > li > strong a, .toc-block ul.refs > li > a { font-weight: 700; }

.footnote { float: footnote; font-size: 7.8pt; line-height: 1.3; text-align: left; hyphens: manual; font-style: normal; font-weight: 400; }
::footnote-call { content: counter(footnote); font-size: 70%; vertical-align: super; line-height: 0; }
::footnote-marker { content: counter(footnote) ". "; }
@page { @footnote { border-top: 0.5pt solid #000; padding-top: 4pt; margin-top: 8pt; } }
"""


def postprocess_html(h):
    # chapter / part headings: split "Chapter 3: Title" into label + title spans
    def split_heading(m):
        attrs, label, title = m.group(1), m.group(2), m.group(3)
        return f'<h1{attrs}><span class="ch-num">{label}</span><span class="ch-title">{title}</span></h1>'
    h = re.sub(r'<h1([^>]*class="(?:chapter|part)"[^>]*)>((?:Chapter|Part) \d+): (.+?)</h1>', split_heading, h)
    def running_head(m):
        attrs, inner = m.group(1), m.group(2)
        title = re.sub(r"<[^>]+>", "", re.sub(r'<span class="ch-num">.*?</span>', "", inner))
        short = re.split(r"(?<=[!?\w”])\s*:\s", title, 1)[0].strip()
        return f'<h1{attrs} data-rh="{html.escape(short, quote=True)}">{inner}</h1>'
    h = re.sub(r"<h1([^>]*)>(.*?)</h1>", running_head, h, flags=re.S)
    # ordered lists that continue numbering (cheat sheet 7., 11., ...): WeasyPrint ignores start=
    h = re.sub(r'<ol start="(\d+)"', lambda m: f'<ol start="{m.group(1)}" style="counter-reset: list-item {int(m.group(1)) - 1}"', h)
    # short closing sections share a page in print
    h = re.sub(r'<h1 id="(whats-next|stay-connected|about-the-author)"', r'<h1 class="runon" id="\1"', h)
    # moment heading + its italic subtitle stay together
    h = re.sub(r'(<h2[^>]*class="moment"[^>]*>.*?</h2>\s*)<p><em>', r'\1<p class="subhead"><em>', h, flags=re.S)
    # paragraphs that are only a bold label ("Say this:") / fully bold scripts / italic stage directions
    h = re.sub(r'<p><strong>([^<]{1,40}:)</strong></p>', r'<p class="label"><strong>\1</strong></p>', h)
    h = re.sub(r'<p><strong>((?:(?!</?strong>|</?p>).)+)</strong></p>', lambda m: m.group(0) if re.search(r"\(Chapter \d+\)$", m.group(1)) else f'<p class="script"><strong>{m.group(1)}</strong></p>', h, flags=re.S)
    h = re.sub(r'<p><em>(\((?:(?!</?p>).)+\))</em></p>', r'<p class="stage"><em>\1</em></p>', h, flags=re.S)
    # footnotes: move each note's text inline so WeasyPrint places it at the foot of the page
    notes = {}
    sec = re.search(r'<section[^>]*class="footnotes[^"]*"[^>]*>.*?</section>', h, re.S)
    if sec:
        for m in re.finditer(r'<li id="(fn\d+)"[^>]*>(.*?)</li>', sec.group(0), re.S):
            txt = re.sub(r'<a href="#fnref\d+"[^>]*>.*?</a>', "", m.group(2), flags=re.S)
            txt = re.sub(r"</?p>", "", txt).strip()
            notes[m.group(1)] = txt
        h = h.replace(sec.group(0), "")
    h = re.sub(r'<a href="#(fn\d+)" class="footnote-ref"[^>]*><sup>\d+</sup></a>',
               lambda m: f'<span class="footnote">{notes[m.group(1)]}</span>', h)
    # contents + find-your-moment lists get page references
    for cls in ("fym", "toc"):
        m = re.search(rf'(<h1[^>]*class="[^"]*{cls}[^"]*"[^>]*>.*?</h1>)(.*?)(?=<h1|<div class="pagebreak")', h, re.S)
        if m:
            block = re.sub(r"<ul>", '<ul class="refs">', m.group(2))
            h = h.replace(m.group(0), m.group(1) + f'<div class="{cls}-block">' + block + "</div>")
    h = re.sub(r'(<h1[^>]*class="part"[^>]*>.*?</h1>\s*<p><em>.*?</em></p>)', r'<div class="partpage">\1</div>', h, flags=re.S)
    h = re.sub(r'(<h1[^>]*class="[^"]*(?:fym|toc)[^"]*"[^>]*>.*?</h1><div class="(?:fym|toc)-block">.*?</div>)',
               r'<div class="frontlist">\1</div>', h, flags=re.S)
    return h


def build_pdf(md, out_path, cfg, work):
    from weasyprint import HTML, CSS
    from weasyprint.text.fonts import FontConfiguration
    fc = FontConfiguration()
    src = os.path.join(work, "print.md")
    open(src, "w", encoding="utf-8").write(md)
    html_path = os.path.join(work, "print.html")
    subprocess.run([pandoc_bin(), src, "-f", "markdown+smart", "-t", "html5", "-s", "--wrap=none",
                    "-o", html_path, "--metadata", f"title={cfg['title']}", "--metadata", f"lang={cfg['language']}"], check=True)
    h = open(html_path, encoding="utf-8").read()
    h = re.sub(r'<header id="title-block-header">.*?</header>', "", h, flags=re.S)
    h = re.sub(r"<style>.*?</style>", "", h, flags=re.S)
    h = postprocess_html(h)
    h = h.replace("<html ", '<html lang="en" ', 1) if 'lang=' not in h[:300] else h
    open(html_path, "w", encoding="utf-8").write(h)
    p = cfg["print"]
    css = (PRINT_CSS.replace("{F}", "file://" + os.path.join(HERE, "fonts"))
           .replace("{W}", str(p["trim_in"][0])).replace("{H}", str(p["trim_in"][1]))
           .replace("{MT}", str(p["margin_top_in"])).replace("{MB}", str(p["margin_bottom_in"]))
           .replace("{MI}", str(p["margin_inside_in"])).replace("{MO}", str(p["margin_outside_in"]))
           .replace("{TITLE}", cfg["title"]))
    open(os.path.join(work, "print.css"), "w").write(css)
    doc = HTML(filename=html_path).render(stylesheets=[CSS(string=css, font_config=fc)], font_config=fc)
    doc.metadata.title = cfg["title"]
    doc.metadata.authors = [cfg["author"]]
    doc.write_pdf(out_path)
    return len(doc.pages)


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("book_dir")
    ap.add_argument("--out")
    ap.add_argument("--keep-work", action="store_true")
    a = ap.parse_args()
    cfg = json.load(open(os.path.join(a.book_dir, "book.json")))
    out = a.out or os.path.join(a.book_dir, "release", "v" + cfg["version"])
    os.makedirs(out, exist_ok=True)
    work = os.path.join(out, "_work") if a.keep_work else tempfile.mkdtemp()
    os.makedirs(work, exist_ok=True)
    base = cfg["title"].replace(" ", "-")
    docx_path = os.path.join(out, f"{base}_Kindle.docx")
    pdf_path = os.path.join(out, f"{base}_Paperback-Interior_{cfg['print']['trim_in'][0]}x{cfg['print']['trim_in'][1]}.pdf")
    build_docx(assemble(a.book_dir, cfg, "kindle"), docx_path, cfg, work)
    pages = build_pdf(assemble(a.book_dir, cfg, "print"), pdf_path, cfg, work)
    print(f"DOCX  {docx_path}\nPDF   {pdf_path}  ({pages} pages)")


if __name__ == "__main__":
    main()
